import json
import unicodedata
from datetime import datetime
from pathlib import Path
import pandas as pd
import requests
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from services.address_service import AddressService


class RBMLQService:
    def __init__(self, url: str, municipios: list[str]):
        self.url = url
        self.municipios_norm = {self._remover_acentos(m) for m in municipios}

    @staticmethod
    def _remover_acentos(texto: str) -> str:
        if not isinstance(texto, str):
            return ""
        texto = unicodedata.normalize("NFKD", texto)
        return "".join(c for c in texto if not unicodedata.combining(c)).strip().upper()

    @staticmethod
    def _sanitizar_df(df: pd.DataFrame) -> pd.DataFrame:
        clean = lambda v: ILLEGAL_CHARACTERS_RE.sub("", str(v)) if pd.notnull(v) and isinstance(v, str) else v
        return df.map(clean) if hasattr(df, "map") else df.applymap(clean)

    def baixar_dados(self, save_raw_path: Path | None = None) -> list[dict]:
        headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"}
        res = requests.get(self.url, headers=headers, timeout=60)
        res.raise_for_status()
        dados = res.json()

        if save_raw_path:
            save_raw_path = Path(save_raw_path)
            save_raw_path.parent.mkdir(parents=True, exist_ok=True)
            with open(save_raw_path, "w", encoding="utf-8") as f:
                json.dump(dados, f, ensure_ascii=False, indent=2)

        if isinstance(dados, dict):
            chave_lista = next((k for k, v in dados.items() if isinstance(v, list)), None)
            return dados[chave_lista] if chave_lista else [dados]
        return dados

    def processar_e_estruturar(self, itens: list[dict]) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
        instrumentos, faixas_rows, historico_rows = [], [], []
        registro_id = 1

        for item in itens:
            if not isinstance(item, dict):
                continue

            muni_raw = item.get("Municipio") or item.get("municipio") or ""
            if self._remover_acentos(muni_raw) not in self.municipios_norm:
                continue

            faixas = item.get("Faixas") or []
            historico = item.get("Historico") or []

            series_list = [f.get("NumeroSerie") for f in faixas if isinstance(f, dict) and f.get("NumeroSerie")]
            series_unicas = list(dict.fromkeys(series_list))
            num_serie = " / ".join(series_unicas) if series_unicas else f"S_SERIE_{registro_id}"
            registro_id += 1

            inmetro_list = [str(f.get("NumeroInmetro")) for f in faixas if isinstance(f, dict) and f.get("NumeroInmetro")]
            inmetro_unicos = list(dict.fromkeys(inmetro_list))
            num_inmetro = " / ".join(inmetro_unicos) if inmetro_unicos else ""

            proprietario = item.get("Proprietario") or {}
            nome_proprietario = (
                proprietario.get("Nome") if isinstance(proprietario, dict) else item.get("Proprietario.Nome")
            )

            # Linha consolidada Instrumentos
            instrumentos.append({
                "NumeroSerie": num_serie,
                "NumeroInmetro": num_inmetro,
                "Proprietario.Nome": nome_proprietario,
                "LocalVerificacao": item.get("LocalVerificacao"),
                "Municipio": item.get("Municipio"),
                "TipoMedidor": item.get("TipoMedidor"),
                "UltimoResultado": item.get("UltimoResultado"),
                "DataUltimaVerificacao": item.get("DataUltimaVerificacao"),
                "DataValidade": item.get("DataValidade"),
                "tipo_local": None,
                "logradouro_limpo": None,
                "numero_limpo": None,
                "rodovia_nome": None,
                "km_num": None,
                "local_maps": None,
                "UF": None,
                "latitude": None,
                "longitude": None
            })

            for f in faixas:
                if isinstance(f, dict):
                    faixas_rows.append({
                        "NumeroSerieInstrumento": num_serie,
                        "NumeroFaixa": f.get("NumeroFaixa"),
                        "NumeroInmetro": f.get("NumeroInmetro"),
                        "NumeroSerieFaixa": f.get("NumeroSerie"),
                        "Sentido": f.get("Sentido"),
                        "VelocidadeNominal": f.get("VelocidadeNominal")
                    })

            for h in historico:
                if isinstance(h, dict):
                    historico_rows.append({
                        "NumeroSerieInstrumento": num_serie,
                        "Ano": h.get("Ano"),
                        "DataLaudo": h.get("DataLaudo"),
                        "DataValidade": h.get("DataValidade"),
                        "TipoServico": h.get("TipoServico"),
                        "Resultado": h.get("Resultado"),
                        "NumeroEnsaio": h.get("NumeroEnsaio"),
                        "NumeroCertificado": h.get("NumeroCertificado")
                    })

        df_inst = self._sanitizar_df(pd.DataFrame(instrumentos))
        df_faixas = self._sanitizar_df(pd.DataFrame(faixas_rows))
        df_hist = self._sanitizar_df(pd.DataFrame(historico_rows))

        return df_inst, df_faixas, df_hist

    def salvar_excel(self, df_inst: pd.DataFrame, df_faixas: pd.DataFrame, df_hist: pd.DataFrame, out_path: Path | str):
        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)

        agora_str = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
        df_controle = pd.DataFrame([
            {"Parametro": "Ultima Atualizacao", "Valor": agora_str},
            {"Parametro": "Total de Instrumentos", "Valor": len(df_inst)},
            {"Parametro": "Total de Faixas Monitoradas", "Valor": len(df_faixas)},
            {"Parametro": "Total de Historicos de Verificacao", "Valor": len(df_hist)},
            {"Parametro": "Fonte dos Dados", "Valor": "RBMLQ / PSIE - Inmetro"},
            {"Parametro": "Status", "Valor": "Atualizado com Sucesso"}
        ])

        with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
            df_controle.to_excel(writer, sheet_name="Controle", index=False)
            df_inst.to_excel(writer, sheet_name="Instrumentos", index=False)
            df_faixas.to_excel(writer, sheet_name="Faixas", index=False)
            df_hist.to_excel(writer, sheet_name="Historico", index=False)

            wb = writer.book

            estilo_tabela = TableStyleInfo(
                name="TableStyleMedium9",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False
            )

            for ws in wb.worksheets:
                if ws.max_row > 1:
                    max_col_letter = get_column_letter(ws.max_column)
                    ref = f"A1:{max_col_letter}{ws.max_row}"
                    nome_tab = f"Tab_{ws.title}"
                    tab = Table(displayName=nome_tab, ref=ref)
                    tab.tableStyleInfo = estilo_tabela
                    ws.add_table(tab)

                    for col in ws.columns:
                        col_letter = get_column_letter(col[0].column)
                        max_len = max(len(str(cell.value or "")) for cell in col)
                        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 45)